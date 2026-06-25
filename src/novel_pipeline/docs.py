"""Document loading, saving, draft staging, and structural validation.

Configurable formats (I6/I7/I8/I9 via config):
  living_doc_backup_format        — backup file naming
  rejected_draft_name_format      — staged draft naming
  canonical_chapter_name_format   — final canonical chapter naming

Promotion behaviour (M1):
  If a canonical chapter already exists at the target path, promote_chapter
  now raises PromotionCollisionError rather than writing a silent
  timestamped duplicate that the state-file/regex layer cannot see. The
  caller is responsible for resolving the conflict.
"""

from __future__ import annotations

import difflib
import os
import shutil
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from .exceptions import DocumentLoadError, PromotionCollisionError
from .logging_ import log_event


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _utc_timestamp() -> str:
    """Compact UTC timestamp suitable for filenames: 20260515T142301Z."""
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def _next_non_colliding(parent: Path, render) -> Path:
    """Given a `render(suffix_index: int | None) -> str` callable that
    produces a candidate filename for a given collision-suffix index,
    return the first non-colliding Path under `parent`."""
    candidate = parent / render(None)
    if not candidate.exists():
        return candidate
    i = 1
    while True:
        candidate = parent / render(i)
        if not candidate.exists():
            return candidate
        i += 1


def _format_with_collision_suffix(base_name: str, suffix_index: int | None) -> str:
    """Apply a collision suffix to a rendered filename.

    The suffix is inserted before the file extension (if the name ends in
    a recognisable extension like .md, .json, .txt) and otherwise appended
    to the end. This keeps the file extension stable (so .md files stay
    .md) while ensuring backup-naming patterns like `<name>.bak.<ts>`
    don't lose their `<name>.bak.` prefix.
    """
    if suffix_index is None:
        return base_name
    sep = f"-{suffix_index:02d}"
    # Only treat the trailing dot-segment as an extension if it's short
    # and alphanumeric (so timestamps like ".20260517T151152Z" aren't
    # treated as extensions).
    if "." in base_name:
        stem, _, ext = base_name.rpartition(".")
        if ext.isalnum() and len(ext) <= 5:
            return f"{stem}{sep}.{ext}"
    return f"{base_name}{sep}"


def _atomic_write(path: Path, content: str) -> None:
    """Write `content` to `path` atomically: tmp + fsync + os.replace."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_path = tempfile.mkstemp(
        prefix=path.name + ".",
        suffix=".tmp",
        dir=str(path.parent),
    )
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(content)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp_path, path)
    except Exception:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
        raise


# ---------------------------------------------------------------------------
# Static docs (templates loaded once per session)
# ---------------------------------------------------------------------------

def _load_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise DocumentLoadError(
            f"python-docx required for .docx files: {path}. "
            f"Install it with: pip install python-docx"
        ) from e
    try:
        doc = docx.Document(str(path))
    except Exception as e:
        raise DocumentLoadError(f"failed to open docx {path}: {e}") from e
    return "\n\n".join(p.text for p in doc.paragraphs)


def load_static_docs(paths: list[str]) -> dict[str, str]:
    """Load template files once. Returns {filename_without_ext: content}."""
    docs: dict[str, str] = {}
    for p in paths:
        path = Path(p)
        if not path.exists():
            raise FileNotFoundError(f"static doc not found: {path.resolve()}")

        ext = path.suffix.lower()
        if ext in (".md", ".txt", ".markdown"):
            try:
                content = path.read_text(encoding="utf-8")
            except UnicodeDecodeError as e:
                raise DocumentLoadError(
                    f"utf-8 decode failed for {path}: {e}"
                ) from e
        elif ext == ".docx":
            content = _load_docx(path)
        elif ext == ".pdf":
            raise DocumentLoadError(
                f"PDF not supported in v1: {path}. Convert to .md or .docx."
            )
        else:
            raise DocumentLoadError(
                f"unsupported extension: {ext or '<none>'} at {path}"
            )

        if not content.strip():
            raise DocumentLoadError(f"empty: {path}")

        key = path.stem
        if key in docs:
            raise ValueError(
                f"static doc name collision on stem '{key}' "
                f"(two files would share that key)"
            )
        docs[key] = content
    return docs


# ---------------------------------------------------------------------------
# Living doc
# ---------------------------------------------------------------------------

def load_living_doc(path: str) -> str:
    """Load the mutable living doc. Missing → empty string + warning log."""
    p = Path(path)
    if not p.exists():
        log_event("living_doc_missing_first_run", {"path": str(p)})
        return ""
    try:
        text = p.read_text(encoding="utf-8")
    except UnicodeDecodeError as e:
        raise DocumentLoadError(f"utf-8 decode failed for {p}: {e}") from e
    return text.rstrip()


def save_living_doc(
    path: str,
    content: str,
    config: dict | None = None,
) -> None:
    """Atomic save with indefinite backup retention.

    I7: backup file naming format is configurable via
    `living_doc_backup_format` (placeholders: {name}, {ts}). Default
    "{name}.bak.{ts}".
    """
    if not content or not content.strip():
        raise ValueError("refusing to save empty living doc")

    cfg = config or {}
    backup_fmt = cfg.get("living_doc_backup_format", "{name}.bak.{ts}")

    p = Path(path)

    if p.exists():
        ts = _utc_timestamp()

        def render(suffix_index):
            base = backup_fmt.format(name=p.name, ts=ts)
            return _format_with_collision_suffix(base, suffix_index)

        backup = _next_non_colliding(p.parent, render)
        shutil.copy2(p, backup)
        log_event("living_doc_backup_written", {"backup": str(backup)})

    _atomic_write(p, content)
    log_event("living_doc_saved", {"path": str(p), "bytes": len(content)})


# ---------------------------------------------------------------------------
# Chapter staging (v1 patch 1: draft → .rejected/ → promote on approval)
# ---------------------------------------------------------------------------

def save_chapter_draft(
    output_dir: str,
    chapter_num: int,
    content: str,
    config: dict | None = None,
) -> str:
    """Write a draft chapter to the .rejected/ staging directory.

    I8: filename format is configurable via `rejected_draft_name_format`
    (placeholders: {nn}, {ts}). Default "chapter_{nn:02d}__{ts}.md".

    Returns the path written. Drafts only move to canonical on approval
    via promote_chapter(). Rejected drafts stay here as an audit trail.
    """
    if not content or not content.strip():
        raise ValueError("refusing to save empty chapter draft")

    cfg = config or {}
    name_fmt = cfg.get("rejected_draft_name_format", "chapter_{nn:02d}__{ts}.md")

    rejected = Path(output_dir) / ".rejected"
    rejected.mkdir(parents=True, exist_ok=True)
    ts = _utc_timestamp()

    def render(suffix_index):
        base = name_fmt.format(nn=chapter_num, ts=ts)
        return _format_with_collision_suffix(base, suffix_index)

    draft_path = _next_non_colliding(rejected, render)
    _atomic_write(draft_path, content)
    log_event(
        "chapter_draft_saved",
        {"path": str(draft_path), "chapter": chapter_num, "bytes": len(content)},
    )
    return str(draft_path)


def promote_chapter(
    draft_path: str,
    output_dir: str,
    chapter_num: int,
    config: dict | None = None,
) -> str:
    """Move an approved draft from .rejected/ to canonical chapter_NN.md.

    M1: if the target already exists this is a state-drift bug — raises
    PromotionCollisionError so the caller can abort cleanly instead of
    silently writing a timestamped duplicate that the regex/state layer
    cannot see.

    I9: canonical filename format is configurable via
    `canonical_chapter_name_format` (placeholder: {nn}). Default
    "chapter_{nn:02d}.md".

    Note: caller is responsible for updating .pipeline_state.json after
    this returns.
    """
    src = Path(draft_path)
    if not src.exists():
        raise FileNotFoundError(f"draft not found for promotion: {src}")

    cfg = config or {}
    name_fmt = cfg.get("canonical_chapter_name_format", "chapter_{nn:02d}.md")
    target = Path(output_dir) / name_fmt.format(nn=chapter_num)

    if target.exists():
        log_event(
            "promote_chapter_target_exists",
            {"target": str(target), "chapter": chapter_num},
        )
        raise PromotionCollisionError(
            f"Cannot promote chapter {chapter_num}: target {target} already "
            f"exists. This usually means resume detection failed to identify "
            f"the chapter as already promoted. Inspect the target file and "
            f"either (a) delete it if it is stale and re-run with --resume, "
            f"or (b) keep it and remove the draft at {src} manually."
        )

    target.parent.mkdir(parents=True, exist_ok=True)
    os.replace(src, target)
    log_event(
        "chapter_promoted",
        {"chapter": chapter_num, "from": str(src), "to": str(target)},
    )
    return str(target)


def find_unpromoted_drafts(output_dir: str, chapter_num: int) -> list[Path]:
    """H7: list .rejected/ files matching this chapter, newest first.

    Used on resume to surface unpromoted drafts so the user can recover
    work instead of regenerating from scratch.
    """
    rejected = Path(output_dir) / ".rejected"
    if not rejected.exists():
        return []
    prefix = f"chapter_{chapter_num:02d}__"
    drafts = [p for p in rejected.iterdir() if p.is_file() and p.name.startswith(prefix)]
    # Newest first by mtime.
    drafts.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return drafts


# ---------------------------------------------------------------------------
# Structural validation (v1 patch 3)
# ---------------------------------------------------------------------------

def validate_living_doc_structure(
    content: str,
    required_sections: list[str],
) -> tuple[bool, list[str]]:
    """Check that all required section headers appear, in order.

    Substring match (case-sensitive) on a line basis. Returns
    (ok, missing_or_out_of_order).
    """
    lines = content.splitlines()

    found_at: dict[str, int] = {}
    for section in required_sections:
        for idx, line in enumerate(lines):
            if section in line:
                found_at[section] = idx
                break

    problems: list[str] = []
    last_idx = -1
    for section in required_sections:
        if section not in found_at:
            problems.append(section)
            continue
        if found_at[section] <= last_idx:
            problems.append(section)
        else:
            last_idx = found_at[section]

    return (len(problems) == 0, problems)


def build_living_doc_diff(old: str, new: str, context_lines: int = 3) -> str:
    """Unified diff of old vs new living doc, for surfacing to humans on
    validation failure."""
    diff = difflib.unified_diff(
        old.splitlines(keepends=True),
        new.splitlines(keepends=True),
        fromfile="living_doc.md (previous)",
        tofile="living_doc.md (proposed)",
        n=context_lines,
    )
    return "".join(diff)
